import os
import openpyxl

etdsDir = "\\\\Lincoln\\Library\\ETDs"

def normalize_header(header):
    return header.lower().strip()

def get_normalized_headers(sheet):
    headers = [normalize_header(sheet.cell(row=1, column=col).value) for col in range(1, sheet.max_column + 1)]
    return headers

def get_row_data(sheet, headers, normalized_headers):
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = [None] * len(headers)
        for col_idx, cell_value in enumerate(row):
            header = normalized_headers[col_idx]
            if header in headers:
                row_data[headers.index(header)] = cell_value
        rows.append(row_data)
    return rows

def combine_excel_files(original_file, batch2_file, batch3_file, output_file):
    # Load the original workbook and get the column order
    original_wb = openpyxl.load_workbook(original_file)
    original_ws = original_wb.active

    # Get normalized column headers from original workbook
    original_headers = [normalize_header(original_ws.cell(row=1, column=col).value) for col in range(1, original_ws.max_column + 1)]

    # Prepare the output workbook and worksheet
    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active

    # Write the headers to the new worksheet
    new_ws.append([original_ws.cell(row=1, column=col).value for col in range(1, original_ws.max_column + 1)])

    # Helper function to process each file
    def process_file(file):
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        ws_headers = get_normalized_headers(ws)
        rows = get_row_data(ws, original_headers, ws_headers)
        return rows

    # Combine data from all files
    combined_rows = process_file(original_file) + process_file(batch2_file) + process_file(batch3_file)

    # Write combined data to the new worksheet
    for row in combined_rows:
        new_ws.append(row)

    # Save the new workbook
    new_wb.save(output_file)

# Example usage
original_file = os.path.join(etdsDir, 'embargos_original.xlsx')
batch2_file = os.path.join(etdsDir, 'ETD Embargoed Submissions 2023 to 4-10-24.xlsx')
batch3_file = os.path.join(etdsDir, 'ETD Embargoed Submissions 4-10-24 to 6-3-24.xlsx')
output_file = os.path.join(etdsDir, 'embargos.xlsx')

combine_excel_files(original_file, batch2_file, batch3_file, output_file)
